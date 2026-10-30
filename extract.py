import docx

doc = docx.Document(r'C:\Users\ASUS\.gemini\antigravity\brain\16cac747-24f1-4ca1-936e-3618c2f9fb0a\.user_uploaded\media_1793339788738.docx')

with open(r'd:\Mantis\blueprint_full.txt', 'w', encoding='utf-8') as f:
    for i, p in enumerate(doc.paragraphs):
        style = p.style.name if p.style else 'None'
        f.write(f'[{i}][{style}] {p.text}\n')
    
    # Also get tables
    for ti, table in enumerate(doc.tables):
        f.write(f'\n=== TABLE {ti} ===\n')
        for ri, row in enumerate(table.rows):
            cells = [cell.text for cell in row.cells]
            sep = " | "
            f.write(f'  ROW {ri}: {sep.join(cells)}\n')

print('Done')
